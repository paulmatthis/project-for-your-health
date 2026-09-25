#!/usr/bin/env python3
"""
Parses food-log.md's structured table rows (never the prose "Day total"
notes, which can go stale) and makes the workbook's Food Log tab an exact
mirror of them. Daily Summary and Weekly Rollup are formula-driven off the
Food Log tab already, so once it's synced they recompute themselves - no
hand-written total, no arithmetic done by a person or an LLM.

Usage: python3 app/recompute.py [--dry-run] [--no-commit]
"""

import argparse
import copy
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FOOD_LOG = PROJECT_ROOT / "food-log.md"
WORKBOOK = PROJECT_ROOT / "project-200-tracker.xlsx"
SHEET_NAME = "Food Log"

ROW_RE = re.compile(
    r"^\|\s*(\d{4}-\d{2}-\d{2})\s*\|(.*)\|\s*(-?\d+)\s*\|\s*(-?\d+)\s*\|"
    r"\s*(-?\d+)[^|]*\|\s*(-?\d+)\s*\|\s*(Yes|No)[^|]*\|(.*)\|\s*$"
)


def parse_food_log(path: Path):
    """Return a list of dicts, one per structured table row, in file order.
    Anything that isn't a table row (headers, separators, prose Notes
    blocks) is ignored - this is the whole point: the app never trusts a
    hand-written summary, only the raw rows."""
    rows = []
    skipped = []
    for lineno, line in enumerate(path.read_text().splitlines(), start=1):
        if not line.startswith("|"):
            continue
        if line.startswith("| Date ") or line.startswith("|---"):
            continue
        m = ROW_RE.match(line)
        if not m:
            skipped.append((lineno, line))
            continue
        date, item, cal, protein, carbs, fat, alcohol_word, confidence = m.groups()
        rows.append(
            {
                "date": date,
                "item": item.strip(),
                "calories": int(cal),
                "protein": int(protein),
                "carbs": int(carbs),
                "fat": int(fat),
                "alcohol": line.split("|")[7].strip(),  # keep full text e.g. "No (0.0% ABV)"
                "confidence": confidence.strip(),
                "line": lineno,
            }
        )
    return rows, skipped


def sync_workbook(rows, dry_run=False):
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb[SHEET_NAME]

    # Template formatting from an existing data row (row 2), so new rows
    # match the workbook's existing convention (blue text, number formats).
    template_fonts = {c: copy.copy(ws.cell(row=2, column=c).font) for c in range(1, 9)}
    template_formats = {c: ws.cell(row=2, column=c).number_format for c in range(1, 9)}

    # Wipe everything from row 2 down to the current last used row - data
    # rows AND any interspersed "Notes, ..." text rows. The workbook's job
    # is numbers; narrative notes live in food-log.md only from now on, so
    # they don't need hand-copying into the sheet (that copying was manual
    # busywork with no formula benefit, and another way for drift to creep
    # in).
    last_row = ws.max_row
    for r in range(2, last_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    for i, row in enumerate(rows, start=2):
        values = [
            row["date"],
            row["item"],
            row["calories"],
            row["protein"],
            row["carbs"],
            row["fat"],
            row["alcohol"],
            row["confidence"],
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c)
            if c == 1:
                import datetime as _dt

                y, mo, d = (int(x) for x in v.split("-"))
                cell.value = _dt.datetime(y, mo, d)
            else:
                cell.value = v
            cell.font = template_fonts[c]
            cell.number_format = template_formats[c]

    if dry_run:
        print(f"[dry-run] would write {len(rows)} rows to '{SHEET_NAME}' (rows 2-{len(rows)+1})")
        return False

    wb.save(WORKBOOK)
    print(f"Synced {len(rows)} rows into '{SHEET_NAME}' (rows 2-{len(rows)+1}).")
    return True


def sync_daily_summary_dates(rows, dry_run=False):
    """Fill in the Date column (A) on Daily Summary for any day that has
    Food Log rows but no date entered yet. The B-N formulas are pre-built
    ~220 days ahead and key off column A (IF(A<row>="","",...)), so a day
    with entries but no date here silently shows as blank/"No data
    logged" even though sync_workbook() already wrote its rows into the
    Food Log tab. Handled here automatically instead of relying on a
    person to remember to add the date by hand."""
    import datetime as _dt

    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb["Daily Summary"]

    anchor_row = 5
    anchor_date = ws.cell(row=anchor_row, column=1).value
    if not isinstance(anchor_date, _dt.datetime):
        print(
            f"WARNING: Daily Summary row {anchor_row} has no anchor date; "
            "skipping date sync.",
            file=sys.stderr,
        )
        return False

    template_font = copy.copy(ws.cell(row=anchor_row, column=1).font)
    template_format = ws.cell(row=anchor_row, column=1).number_format

    log_dates = sorted({row["date"] for row in rows})
    added = []
    for date_str in log_dates:
        y, mo, d = (int(x) for x in date_str.split("-"))
        target_date = _dt.datetime(y, mo, d)
        target_row = anchor_row + (target_date - anchor_date).days
        if target_row < anchor_row or target_row > ws.max_row:
            continue  # before the anchor or past the pre-built range
        cell = ws.cell(row=target_row, column=1)
        if cell.value is None:
            cell.value = target_date
            cell.font = template_font
            cell.number_format = template_format
            added.append(date_str)

    if not added:
        return False
    if dry_run:
        print(f"[dry-run] would add {len(added)} missing date(s) to Daily Summary: {', '.join(added)}")
        return False

    wb.save(WORKBOOK)
    print(f"Added {len(added)} missing date(s) to Daily Summary: {', '.join(added)}")
    return True


def recalc():
    import glob
    import os

    env = os.environ.copy()
    env["PATH"] = "/Applications/LibreOffice.app/Contents/MacOS:" + env.get("PATH", "")

    candidates = glob.glob(
        str(Path.home() / "Library/Application Support/Claude/**/skills/xlsx/scripts/recalc.py"),
        recursive=True,
    )
    if not candidates:
        print("Could not locate the xlsx skill's recalc.py; skipping recalculation.", file=sys.stderr)
        return None
    result = subprocess.run(
        [sys.executable, candidates[0], str(WORKBOOK), "90"],
        capture_output=True,
        text=True,
        env=env,
    )
    print(result.stdout.strip())
    if result.returncode != 0:
        print(result.stderr, file=sys.stderr)
    return result


def git_commit_and_push(message: str):
    # Scoped to WORKBOOK specifically (both the diff check and the commit
    # itself) so this never sweeps up an unrelated change that happens to
    # be staged elsewhere in the tree at the same time this script runs.
    subprocess.run(["git", "add", str(WORKBOOK)], cwd=PROJECT_ROOT, check=True)
    diff = subprocess.run(
        ["git", "diff", "--cached", "--quiet", "--", str(WORKBOOK)], cwd=PROJECT_ROOT
    )
    if diff.returncode == 0:
        print("No changes to commit.")
        return
    subprocess.run(["git", "commit", "-m", message, "--", str(WORKBOOK)], cwd=PROJECT_ROOT, check=True)

    push = subprocess.run(["git", "push"], cwd=PROJECT_ROOT, capture_output=True, text=True)
    if push.returncode != 0:
        print("Push rejected, pulling and retrying once...", file=sys.stderr)
        pull = subprocess.run(
            ["git", "pull", "--rebase"], cwd=PROJECT_ROOT, capture_output=True, text=True
        )
        if pull.returncode != 0:
            print(f"git pull --rebase failed:\n{pull.stderr}", file=sys.stderr)
            print(
                "Could not auto-resolve. Stopping here rather than guessing - "
                "check the working tree by hand.",
                file=sys.stderr,
            )
            sys.exit(1)
        retry = subprocess.run(["git", "push"], cwd=PROJECT_ROOT, capture_output=True, text=True)
        if retry.returncode != 0:
            print(f"Push still failing after retry:\n{retry.stderr}", file=sys.stderr)
            sys.exit(1)
    print("Committed and pushed.")


def print_summary(rows):
    from collections import defaultdict

    by_date = defaultdict(lambda: {"calories": 0, "protein": 0, "carbs": 0, "fat": 0})
    for row in rows:
        d = by_date[row["date"]]
        d["calories"] += row["calories"]
        d["protein"] += row["protein"]
        d["carbs"] += row["carbs"]
        d["fat"] += row["fat"]
    if not rows:
        return
    last_date = rows[-1]["date"]
    t = by_date[last_date]
    print(
        f"Latest day in food-log.md ({last_date}): "
        f"{t['calories']} kcal, {t['protein']}g protein, {t['carbs']}g carb, {t['fat']}g fat."
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-commit", action="store_true")
    args = parser.parse_args()

    rows, skipped = parse_food_log(FOOD_LOG)
    if skipped:
        print(f"WARNING: {len(skipped)} table-looking line(s) failed to parse:", file=sys.stderr)
        for lineno, line in skipped:
            print(f"  line {lineno}: {line}", file=sys.stderr)

    print_summary(rows)
    changed = sync_workbook(rows, dry_run=args.dry_run)
    dates_changed = sync_daily_summary_dates(rows, dry_run=args.dry_run)
    if args.dry_run:
        return
    if changed or dates_changed:
        recalc()
        if not args.no_commit:
            git_commit_and_push(
                f"auto: sync Food Log tab from food-log.md ({len(rows)} rows)"
            )


if __name__ == "__main__":
    main()
